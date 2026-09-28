"""
Tests for vacation percent tiers, EmplDetl parser, running balance helpers, sick pay hours.
"""

from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

from services.vacation_percent_service import (
    resolve_vacation_percent_from_tiers,
    years_of_service,
    to_hourly_rate,
    BC_ESA_DEFAULT_TIERS,
)
from services.payroll_detail_import_service import parse_empldetl_workbook, _normalize_name, match_employee_by_name
from services.leave_pay_helpers import leave_days_overlapping


class FakeEmp:
    def __init__(self, id, full_name, first_name=None, last_name=None):
        self.id = id
        self.full_name = full_name
        self.first_name = first_name
        self.last_name = last_name


def test_years_of_service():
    assert abs(years_of_service(date(2020, 1, 1), date(2025, 1, 1)) - 5.0) < 0.02
    assert years_of_service(None, date(2025, 1, 1)) == 0.0


def test_bc_esa_tier_resolution():
    tiers = [
        {"min_years": mn, "max_years": mx, "percent": pct}
        for mn, mx, pct in BC_ESA_DEFAULT_TIERS
    ]
    assert resolve_vacation_percent_from_tiers(tiers, 0) == Decimal("4.0")
    assert resolve_vacation_percent_from_tiers(tiers, 4.9) == Decimal("4.0")
    assert resolve_vacation_percent_from_tiers(tiers, 5.0) == Decimal("6.0")
    assert resolve_vacation_percent_from_tiers(tiers, 20.0) == Decimal("6.0")


def test_ca_style_tiers():
    tiers = [
        {"min_years": 0, "max_years": 5, "percent": 4},
        {"min_years": 5, "max_years": 11, "percent": 6},
        {"min_years": 11, "max_years": 21, "percent": 8},
        {"min_years": 21, "max_years": None, "percent": 10},
    ]
    assert resolve_vacation_percent_from_tiers(tiers, 3) == Decimal("4")
    assert resolve_vacation_percent_from_tiers(tiers, 5) == Decimal("6")
    assert resolve_vacation_percent_from_tiers(tiers, 10.87) == Decimal("6")
    assert resolve_vacation_percent_from_tiers(tiers, 11.0) == Decimal("8")
    assert resolve_vacation_percent_from_tiers(tiers, 15) == Decimal("8")
    assert resolve_vacation_percent_from_tiers(tiers, 25) == Decimal("10")


def test_tier_gap_does_not_jump_to_highest():
    """User-entered max 4/10/20 (exclusive) creates gaps; must not fall through to 10%."""
    tiers = [
        {"min_years": 0, "max_years": 4, "percent": 4},
        {"min_years": 5, "max_years": 10, "percent": 6},
        {"min_years": 11, "max_years": 20, "percent": 8},
        {"min_years": 21, "max_years": None, "percent": 10},
    ]
    # ~10.87 years (Harjit Jul 2026, hire Sep 2015) sits in gap [10, 11)
    assert resolve_vacation_percent_from_tiers(tiers, 10.87) == Decimal("6")
    assert resolve_vacation_percent_from_tiers(tiers, 11.0) == Decimal("8")
    assert resolve_vacation_percent_from_tiers(tiers, 4.5) == Decimal("4")


def test_harjit_anniversary_sep_2026():
    hire = date(2015, 9, 1)
    tiers = [
        {"min_years": 0, "max_years": 5, "percent": 4},
        {"min_years": 5, "max_years": 11, "percent": 6},
        {"min_years": 11, "max_years": 21, "percent": 8},
        {"min_years": 21, "max_years": None, "percent": 10},
    ]
    y_before = years_of_service(hire, date(2026, 8, 31))
    y_on = years_of_service(hire, date(2026, 9, 1))
    assert y_before < 11.0
    assert y_on >= 11.0
    assert resolve_vacation_percent_from_tiers(tiers, y_before) == Decimal("6")
    assert resolve_vacation_percent_from_tiers(tiers, y_on) == Decimal("8")


def test_union_and_non_union_schedules_are_independent():
    """Same YOS can resolve differently when union vs non-union schedules differ."""
    non_union = [
        {"min_years": 0, "max_years": 5, "percent": 4},
        {"min_years": 5, "max_years": None, "percent": 6},
    ]
    union = [
        {"min_years": 0, "max_years": 5, "percent": 4},
        {"min_years": 5, "max_years": 11, "percent": 6},
        {"min_years": 11, "max_years": None, "percent": 8},
    ]
    years = 12.0
    assert resolve_vacation_percent_from_tiers(non_union, years) == Decimal("6")
    assert resolve_vacation_percent_from_tiers(union, years) == Decimal("8")


def test_to_hourly_rate():
    assert to_hourly_rate(25.0, "Hourly") == 25.0
    assert abs(to_hourly_rate(2080 * 20, "Annual") - 20.0) < 0.001
    monthly = (52 * 40 / 12) * 30
    assert abs(to_hourly_rate(monthly, "Monthly") - 30.0) < 0.01


def test_running_vacation_balance_math():
    opening = Decimal("1000.00")
    earned = [Decimal("50"), Decimal("60")]
    paid = [Decimal("0"), Decimal("200")]
    bal = opening
    for e, p in zip(earned, paid):
        bal = bal + e - p
    assert bal == Decimal("910.00")


def test_leave_days_overlapping_full_and_partial():
    days = leave_days_overlapping(
        date(2026, 1, 5), date(2026, 1, 9), 5.0,
        date(2026, 1, 1), date(2026, 1, 31),
    )
    assert days == 5.0
    partial = leave_days_overlapping(
        date(2026, 1, 28), date(2026, 2, 5), 7.0,
        date(2026, 1, 1), date(2026, 1, 31),
    )
    assert partial > 0
    assert partial < 7.0


def test_normalize_and_match_name():
    index = {
        _normalize_name("Amandeep Boparai"): [FakeEmp("E1", "Amandeep Boparai", "Amandeep", "Boparai")],
        _normalize_name("Andrew Weatherbee"): [FakeEmp("E2", "Andrew Weatherbee", "Andrew", "Weatherbee")],
    }
    emp, status = match_employee_by_name("AMANDEEP BOPARAI", index)
    assert status == "matched"
    assert emp.id == "E1"
    emp2, status2 = match_employee_by_name("Unknown Person", index)
    assert status2 == "unmatched"
    assert emp2 is None


def _build_sample_empldetl() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Quadra Wood Products Ltd"
    ws["A2"] = "Employee Detail 2026-09-15 to 2026-09-15"
    headers = [None, "Date", "Cheque No.", "Gross", "Withheld", "Net Pay", "EI", "CPP", "CPP2", "Tax",
               "Vacation Paid", "Salary", "Regular Hrs.", "OT Rate 1 Hrs", "STAT Pay Hours", "Meal",
               "Union $0.02 Hours", "Union $0.02", "Union 1.45%", "Union Initia", "Vacation Earned",
               "EI Ins. Earnings", "Benefits"]
    for i, h in enumerate(headers, start=1):
        ws.cell(4, i, h)
    ws["A5"] = "AMANDEEP BOPARAI"
    ws.cell(6, 2, date(2026, 9, 15))
    ws.cell(6, 3, "QPR1-20260915")
    ws.cell(6, 4, 5031.5)   # Gross
    ws.cell(6, 11, 0)       # Vacation Paid
    ws.cell(6, 21, 300)     # Vacation Earned = (5031.5 - 31.5) * 6%
    ws.cell(6, 23, 31.5)    # Benefits
    ws["A7"] = "Total"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_empldetl_sample_shape():
    data = parse_empldetl_workbook(_build_sample_empldetl())
    assert data["company_header"] == "Quadra Wood Products Ltd"
    assert data["period_start"] == "2026-09-15"
    assert len(data["rows"]) == 1
    row = data["rows"][0]
    assert row["employee_name"] == "AMANDEEP BOPARAI"
    assert row["gross"] == 5031.5
    assert row["benefits"] == 31.5
    assert row["vacation_earned"] == 300
    assert abs((row["gross"] - row["benefits"]) * 0.06 - row["vacation_earned"]) < 0.01
    assert row["vacation_paid"] == 0
    assert row["cheque_no"] == "QPR1-20260915"
    assert data["report_kind"] == "all_staff_single_pay"


def _build_single_employee_multi_pay() -> bytes:
    """EmplDetl_Harjit-style: one employee, many cheque dates, YTD title range."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Quadra Wood Products Ltd"
    ws["A2"] = "Employee Detail 2026-01-01 to 2026-09-17"
    headers = [
        None, "Date", "Cheque No.", "Gross", "Withheld", "Net Pay", "EI", "CPP", "CPP2", "Tax",
        "Vacation Paid", "Regular Hrs.", "OT Rate 1 Hrs", "STAT Pay Hours", "Meal",
        "Union $0.02 Hours", "Union $0.02", "Union 1.45%", "Union Initia",
        "Vacation Adj Hours", "Vacation Earned",
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(4, i, h)
    ws["A5"] = "HARJIT GREWAL"
    # Two pay lines
    ws.cell(6, 2, date(2026, 1, 15))
    ws.cell(6, 3, "QPR6-20260115")
    ws.cell(6, 4, 1302)
    ws.cell(6, 11, 0)
    ws.cell(6, 21, 78.12)
    ws.cell(7, 2, date(2026, 9, 15))
    ws.cell(7, 3, "QPR6-20260915")
    ws.cell(7, 4, 2965.49)
    ws.cell(7, 11, 0)
    ws.cell(7, 21, 235)
    ws["A8"] = "Total"
    ws["A9"] = "Generated On: 2026-09-24"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_single_employee_multi_pay():
    data = parse_empldetl_workbook(_build_single_employee_multi_pay())
    assert data["report_kind"] == "single_employee_multi_pay"
    assert len(data["rows"]) == 2
    assert data["rows"][0]["employee_name"] == "HARJIT GREWAL"
    assert data["rows"][0]["gross"] == 1302
    assert data["rows"][0]["vacation_earned"] == 78.12
    # YTD title must not stamp period onto each cheque row
    assert data["rows"][0]["period_start"] is None
    assert data["rows"][1]["pay_date"] == "2026-09-15"
    assert data["rows"][1]["vacation_earned"] == 235
    # "Generated On" must not become an employee name with phantom rows
    assert all(r["employee_name"] == "HARJIT GREWAL" for r in data["rows"])


if __name__ == "__main__":
    test_years_of_service()
    test_bc_esa_tier_resolution()
    test_ca_style_tiers()
    test_tier_gap_does_not_jump_to_highest()
    test_harjit_anniversary_sep_2026()
    test_union_and_non_union_schedules_are_independent()
    test_to_hourly_rate()
    test_running_vacation_balance_math()
    test_leave_days_overlapping_full_and_partial()
    test_normalize_and_match_name()
    test_parse_empldetl_sample_shape()
    test_parse_single_employee_multi_pay()
    print("ALL PASSED")
